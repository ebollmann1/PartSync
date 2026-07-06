from odoo import models, fields, api
import base64
import tempfile
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from odoo.exceptions import UserError
from psycopg2.extras import execute_values, Json
import logging
_logger = logging.getLogger(__name__)

class UpdateProductAttribute(models.TransientModel):
    _name = 'update.product.attribute'
    _description = 'Update product Attribute'

    file = fields.Binary(required=True, string="XLSX File")
    filename = fields.Char()

    def action_update(self):

        file_content = base64.b64decode(self.file)
        workbook = load_workbook(BytesIO(file_content), read_only=True)
        sheet = workbook.worksheets[0]

        headers = [h.value for h in sheet[1]]
        col = {name: i for i, name in enumerate(headers)}
        rows = sheet.iter_rows(min_row=2, values_only=True)
        
        product_template_dict = {}
        product_data = self.env['product.template'].search_read([], fields=['default_code', 'id'])
        for pr in product_data:
            if pr.get('default_code') not in product_template_dict:
                product_template_dict[pr.get('default_code')] = pr.get('id')
        print("\n +++++++++  product_template_dict +++++++", product_template_dict)
        
        attribute_dict = {}
        attribute_data = self.env['product.attribute'].search_read([], fields=['name', 'id'])
        for ad in attribute_data:
            if ad.get('name') not in attribute_dict:
                attribute_dict[ad.get('name')] = ad.get('id')
        print("\n +++++++++  attribute_dict +++++++", attribute_dict)
        
        attribute_value_dict = {}
        attribute_value_data = self.env['product.attribute.value'].search_read([], fields=['name', 'attribute_id', 'id'])
        for avd in attribute_value_data:
            key = str(avd.get('attribute_id')[0]) + '-' + avd.get('name')
            if key not in attribute_value_dict:
                attribute_value_dict[key] = avd.get('id')
        print("\n +++++++++  attribute_value_dict +++++++", attribute_value_dict)
        
        attribute_line_dict = {}
        attribute_line_data = self.env['product.template.attribute.line'].search_read(
            [], fields=['product_tmpl_id', 'attribute_id', 'id']
        )
        for ald in attribute_line_data:
            key = str(ald.get('attribute_id')[0]) + '-' + str(ald.get('product_tmpl_id')[0])
            if key not in attribute_line_dict:
                attribute_line_dict[key] = ald.get('id')
        print("\n +++++++++  attribute_line_dict +++++++", attribute_line_dict)
        count = 2
        for row in rows:
            internal_ref =row[col['Internal Reference']]
            attribute_name = row[col['Attribute']]
            value_name = row[col['Attribute Value']]
            if type(internal_ref) == float:
                internal_ref = str(int(internal_ref))
            
            product_id = False
            if internal_ref in product_template_dict:
                product_id = product_template_dict.get(internal_ref)
            
            print("\n ++++++++++ internal_ref, product_id ++++++", internal_ref, product_id)
            if product_id:
                if type(attribute_name) == float:
                    attribute_name = str(int(attribute_name))
                
                attribute_id = False
                if attribute_name in attribute_dict:
                    attribute_id = attribute_dict.get(attribute_name)
                
                if attribute_id:
                    if type(value_name) == float:
                        value_name = str(int(value_name))
                    
                    key = str(attribute_id) + '-' + value_name
                    value_id = False
                    if key in attribute_value_dict:
                        value_id = attribute_value_dict.get(key)
                    
                    if value_id and attribute_id:
                        key = str(attribute_id) + '-' + str(product_id)
                        attribute_line_id = False
                        if key in attribute_line_dict:
                            attribute_line_id = attribute_line_dict[key]
                        
                        if attribute_line_id:
                            self.env['product.template.attribute.line'].browse(attribute_line_id).write({
                                'value_ids': [(4, value_id)]
                            })
                        else:
                            line_id = self.env['product.template.attribute.line'].create({
                                'product_tmpl_id': product_id,
                                'attribute_id': attribute_id,
                                'value_ids': [(6, 0, [value_id])]
                            })
                            attribute_line_dict[str(attribute_id) + '-' + str(product_id)] = line_id
                print("\n ++++++ count +++++", count)
                _logger.info("______Count _____%s", count)
                count += 1
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': 'Update Attribute',
                'message': (
                    f'Update Attribute Successfully!\n'
                ),
                'type': 'success',
                'sticky': False,
                'next': {'type': 'ir.actions.act_window_close'},
            }
        }